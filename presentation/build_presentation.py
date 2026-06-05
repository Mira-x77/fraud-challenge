"""Generate FraudShield hackathon PowerPoint."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "FraudShield_Presentation.pptx"

DARK = RGBColor(0x08, 0x08, 0x08)
GREEN = RGBColor(0x00, 0xFF, 0x88)
RED = RGBColor(0xFF, 0x28, 0x44)
WHITE = RGBColor(0xF0, 0xF0, 0xF0)
GRAY = RGBColor(0xAA, 0xAA, 0xAA)


def _style_title(shape, size=32, color=GREEN):
    shape.text_frame.paragraphs[0].font.size = Pt(size)
    shape.text_frame.paragraphs[0].font.bold = True
    shape.text_frame.paragraphs[0].font.color.rgb = color
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def _add_bullets(text_frame, lines, size=16, color=WHITE):
    text_frame.clear()
    for i, line in enumerate(lines):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(6)


def _slide_title_content(prs, title, bullets, image=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK

    if image and Path(image).exists():
        slide.shapes.add_picture(str(image), Inches(5.2), Inches(1.0), width=Inches(4.5))
        title_w = Inches(4.8)
    else:
        title_w = Inches(9)

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), title_w, Inches(0.8))
    tb.text_frame.text = title
    _style_title(tb, size=26, color=RED)

    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(4.6 if image else 9), Inches(5.5))
    _add_bullets(body.text_frame, bullets, size=14)
    return slide


def _slide_full_image(prs, title, subtitle, image):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK

    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.2), Inches(0.6))
    tb.text_frame.text = title
    _style_title(tb, size=22, color=GREEN)

    if subtitle:
        st = slide.shapes.add_textbox(Inches(0.4), Inches(0.75), Inches(9.2), Inches(0.4))
        st.text_frame.text = subtitle
        st.text_frame.paragraphs[0].font.size = Pt(12)
        st.text_frame.paragraphs[0].font.color.rgb = GRAY

    if Path(image).exists():
        slide.shapes.add_picture(str(image), Inches(0.4), Inches(1.1), width=Inches(9.2))


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 1 — Title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = DARK
    t = s.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9), Inches(1.2))
    t.text_frame.text = "FraudShield"
    _style_title(t, size=44, color=GREEN)
    sub = s.shapes.add_textbox(Inches(0.5), Inches(3.3), Inches(9), Inches(1.5))
    _add_bullets(
        sub.text_frame,
        [
            "Hackathon INTELO2026 — Détection de fraude financière",
            "Mira-x77 · LBS × Cursor",
            "11/11 tests publics · CI verte",
        ],
        size=18,
        color=WHITE,
    )

    # 2 — Mission
    _slide_title_content(
        prs,
        "Mission",
        [
            "Analyser des transactions bancaires en lot",
            "Détecter la fraude sans bloquer les clients honnêtes",
            "Expliquer chaque alerte (score + raison)",
            "Deux livrables :",
            "  • fraud_detection.py → noté par robot (CI)",
            "  • app.py → démonstration jury (interface HUD)",
        ],
    )

    # 3 — Architecture
    _slide_title_content(
        prs,
        "Architecture technique",
        [
            "fraud_detection.py",
            "  • detect_fraud(transactions) — fonction pure",
            "  • 7 vecteurs de fraude + fusion de scores",
            "  • Jamais de crash sur données imparfaites",
            "",
            "app.py (Streamlit)",
            "  • render_interface() — dashboard broadcast HUD",
            "  • 3 écrans : accueil → chargement → résultats",
            "  • CSS injecté + st.html pour rendu fiable",
        ],
    )

    # 4 — Detection logic
    _slide_title_content(
        prs,
        "Moteur de détection — 7 signaux",
        [
            "1. CNP — carte absente + montant élevé",
            "2. Vélocité — >5 txns en 60 secondes",
            "3. Géo — impossibilité (Haversine / 900 km/h)",
            "4. ATO — nouveau commerçant + montant suspect",
            "5. Remboursements — 3+ montants négatifs / 24h",
            "6. Structuration — montants sous seuils réglementaires",
            "7. Doublons — même paiement en < 5 min",
            "",
            "Score : min(1.0, max(signaux) + 0.05×(N−1))",
            "Alerte si score ≥ 0.5",
        ],
    )

    # 5 — Idle screen
    _slide_full_image(
        prs,
        "Écran 1 — Accueil (idle)",
        "Globe central · bouton ANALYSER · ticker binaire live",
        SHOTS / "01_idle.png",
    )
    _slide_title_content(
        prs,
        "Écran 1 — Ce qu'il montre",
        [
            "État d'attente avant analyse",
            "Globe 3D animé + orbites radar (CSS @keyframes)",
            "Ticker binaire défilant = effet « données live »",
            "Bouton ANALYSER au-dessus du globe",
            "",
            "Comment : app.py + assets/globe.png + CSS injecté",
        ],
        image=SHOTS / "01_idle.png",
    )

    # 6 — Loading
    _slide_title_content(
        prs,
        "Écran 2 — Chargement (~3 sec)",
        [
            "Animation cyberpunk plein écran",
            "Anneaux rotatifs + texte glitch",
            "Barre de progression + flux binaire",
            "Pendant ce temps : detect_fraud() s'exécute",
            "",
            "Implémentation :",
            "  • session_state.phase = loading",
            "  • HTML/CSS animations (pas de GIF externe)",
            "  • time.sleep(2.8) puis passage à results",
        ],
    )

    # 7 — Results dashboard
    _slide_full_image(
        prs,
        "Écran 3 — Résultats broadcast",
        "KPIs · globe · fil d'alertes · graphiques HUD",
        SHOTS / "03_results_overview.png",
    )
    _slide_title_content(
        prs,
        "Écran 3 — Ce qu'il montre",
        [
            "4 KPIs : transactions, alertes, sûres, capital protégé",
            "Fil d'alertes COMPLET (T-004 + T-020 avec raisons)",
            "Graphiques HUD intégrés (pas de charts Streamlit bleus)",
            "Globe « GLOBAL SCAN » + ticker live",
            "",
            "Thème vert = mode résultats (broadcast monitoring actif)",
        ],
        image=SHOTS / "02_results_dashboard.png",
    )

    # 8 — Alerts explained
    _slide_title_content(
        prs,
        "Les 2 fraudes détectées (exemple)",
        [
            "T-004 · score 1.0 · CARTE BLOQUÉE",
            "  • Client U1 dépense ~50€ habituellement",
            "  • 4800€ en ligne, carte absente, nouveau commerçant",
            "  • CNP + montant anormal + ATO combinés",
            "",
            "T-020 · score 0.95 · CARTE BLOQUÉE",
            "  • Montant négatif (-30€)",
            "  • Règle niveau 1 : donnée invalide / remboursement suspect",
        ],
    )

    # 9 — Investigator
    _slide_full_image(
        prs,
        "Panneau enquêteur",
        "Détails complets · historique client · expanders T-004 / T-020",
        SHOTS / "04_investigator.png",
    )
    _slide_title_content(
        prs,
        "Enquêteur — comment c'est fait",
        [
            "Table HUD custom (HTML) — raisons non tronquées",
            "Expanders par transaction signalée",
            "Détails clé/valeur avec montants négatifs en rouge",
            "Historique complet du client dans le lot",
            "",
            "Objectif jury : comprendre POURQUOI sans lire le code",
        ],
        image=SHOTS / "04_investigator.png",
    )

    # 10 — Results & demo
    _slide_title_content(
        prs,
        "Résultats & démo",
        [
            "✅ 11/11 tests publics (CI verte)",
            "✅ PR soumise : Mira-x77/fraud-challenge",
            "✅ Interface full-screen broadcast HUD",
            "",
            "Lancer la démo :",
            "  pip install -r requirements.txt",
            "  streamlit run app.py",
            "",
            "Merci · Questions ?",
        ],
    )

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
